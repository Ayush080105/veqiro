"""Rex reads every row of an upload, not the stored 500-row preview."""
import asyncio
import csv
import io

import pandas as pd
import pytest
from openpyxl import Workbook

from agents.rex import dataset_sql as d

TYPES = {"Date": "date", "City": "categorical", "Amount": "numeric"}


def _rows(n):
    return [{"Date": f"2026-0{1 + i % 3}-1{i % 9}", "City": "Delhi" if i < 500 else "Mumbai",
             "Amount": "₹1,000"} for i in range(n)]


def _csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0]))
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode()


def _preview(rows, n=500):
    return {"headers": list(rows[0]), "columnTypes": TYPES, "rows": rows[:n]}


def _count(ds, sql):
    return asyncio.run(d.run(ds, sql)).rows


def test_csv_reads_every_row_beyond_the_preview():
    rows = _rows(2000)
    ds = d.load_file(_csv(rows), "rex-dataset/org-orders.csv", _preview(rows))
    assert ds.whole_file and ds.total_rows == 2000
    assert _count(ds, 'SELECT "City", count(*) FROM data GROUP BY 1 ORDER BY 1') == [["Delhi", 500], ["Mumbai", 1500]]
    assert _count(ds, 'SELECT sum("Amount") FROM data') == [[2_000_000.0]]


def test_xlsx_finds_the_header_under_a_title_block_and_reads_every_sheet():
    rows = _rows(800)
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["Company — export"])
    ws.append([])
    ws.append(list(rows[0]))
    for r in rows:
        ws.append(list(r.values()))
    t = wb.create_sheet("Targets")
    t.append(["City", "Target"])
    t.append(["Mumbai", "₹4,00,000"])
    buf = io.BytesIO()
    wb.save(buf)
    preview = {**_preview(rows), "sheets": {
        "Orders": _preview(rows),
        "Targets": {"headers": ["City", "Target"], "columnTypes": {"City": "categorical", "Target": "numeric"},
                    "rows": [{"City": "Mumbai", "Target": "₹4,00,000"}]}}}
    ds = d.load_file(buf.getvalue(), "rex-dataset/org-orders.xlsx", preview)
    assert set(ds.tables) == {"orders", "targets"} and ds.total_rows == 801
    assert _count(ds, 'SELECT count(*) FROM orders') == [[800]]
    assert _count(ds, 'SELECT "Target" FROM targets') == [[400000.0]]


def test_a_file_that_does_not_match_its_preview_is_refused_not_misread():
    rows = _rows(10)
    other = _csv([{"Something": "else", "Entirely": "1"}])
    with pytest.raises(d.FileUnusable):
        d.load_file(other, "x.csv", _preview(rows))


def test_unsupported_types_are_refused():
    with pytest.raises(d.FileUnusable):
        d.load_file(b"%PDF", "report.pdf", _preview(_rows(3)))


def test_date_order_comes_from_the_column():
    month_first = d.parse_dates(pd.Series(["06/25/2026", "01/02/2026"]))
    day_first = d.parse_dates(pd.Series(["25/06/2026", "01/02/2026"]))
    ambiguous = d.parse_dates(pd.Series(["01/02/2026"]))
    assert [str(x)[:10] for x in month_first] == ["2026-06-25", "2026-01-02"]
    assert [str(x)[:10] for x in day_first] == ["2026-06-25", "2026-02-01"]
    assert str(ambiguous[0])[:10] == "2026-02-01"  # Indian default: day first


def test_answers_admit_when_only_the_preview_was_read():
    rows = _rows(500)
    assert "first 500 rows" in d.coverage_note(d.load(_preview(rows)))
    assert "all 2000 rows" in d.coverage_note(d.load_file(_csv(_rows(2000)), "a.csv", _preview(_rows(2000))))
    assert "all 30 rows" in d.coverage_note(d.load(_preview(_rows(30))))
